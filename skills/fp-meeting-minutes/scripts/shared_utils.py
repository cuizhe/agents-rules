#!/usr/bin/env python3
"""
会议纪要公共工具库。

封装 OpenXML / python-docx 的底层操作，被 generate_minutes.py、fix_docx_format.py
等脚本复用。格式常量与 references/docx_format_spec.md 保持一致。
"""

from __future__ import annotations

from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

# ── 格式常量 ───────────────────────────────────────────────────────
FONT_NAME = "微软雅黑"
FONT_SIZE_BODY = "21"       # 五号 = 10.5pt = 21 half-points
FONT_SIZE_H1 = "28"         # 14pt
FONT_SIZE_H2 = "24"         # 12pt
FONT_SIZE_TITLE = "44"      # 2号 = 22pt

BORDER_SIZE = 4             # 0.5pt
BORDER_COLOR = "000000"

# ── XML 辅助 ───────────────────────────────────────────────────────


def _set_font(run, font_name: str = FONT_NAME, size: str = FONT_SIZE_BODY, bold: bool = False):
    """设置 run 字体，确保 ascii/hAnsi/eastAsia 三个属性同时写入。"""
    rPr = run._r.find(qn("w:rPr"))
    if rPr is None:
        rPr = etree.SubElement(run._r, qn("w:rPr"))

    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = etree.SubElement(rPr, qn("w:rFonts"))
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)

    sz = rPr.find(qn("w:sz"))
    if sz is None:
        sz = etree.SubElement(rPr, qn("w:sz"))
    sz.set(qn("w:val"), size)

    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = etree.SubElement(rPr, qn("w:szCs"))
    szCs.set(qn("w:val"), size)

    if bold:
        b = rPr.find(qn("w:b"))
        if b is None:
            b = etree.SubElement(rPr, qn("w:b"))


def _clear_para_indent(para):
    """清除段落首行缩进和左缩进，覆盖 Normal 样式继承。"""
    pPr = para._p.find(qn("w:pPr"))
    if pPr is None:
        pPr = etree.SubElement(para._p, qn("w:pPr"))

    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = etree.SubElement(pPr, qn("w:ind"))
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    ind.set(qn("w:left"), "0")


def _set_cell_margins(cell, top=0, left=0, bottom=0, right=0):
    """设置单元格边距为 0。"""
    tc = cell._tc
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = etree.SubElement(tc, qn("w:tcPr"))

    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is not None:
        tcPr.remove(tcMar)
    tcMar = etree.SubElement(tcPr, qn("w:tcMar"))

    for side, val in [("top", top), ("left", left), ("bottom", bottom), ("right", right)]:
        m = etree.SubElement(tcMar, qn(f"w:{side}"))
        m.set(qn("w:w"), str(val))
        m.set(qn("w:type"), "dxa")


def _set_table_cell_margins(table):
    """设置表格级别单元格边距为 0。"""
    tblPr = table._tbl.tblPr
    tblCellMar = tblPr.find(qn("w:tblCellMar"))
    if tblCellMar is not None:
        tblPr.remove(tblCellMar)
    tblCellMar = etree.SubElement(tblPr, qn("w:tblCellMar"))

    for side in ["top", "start", "bottom", "end"]:
        m = etree.SubElement(tblCellMar, qn(f"w:{side}"))
        m.set(qn("w:w"), "0")
        m.set(qn("w:type"), "dxa")


def _set_cell_width(cell, width: int):
    """设置单元格宽度（DXA）。"""
    tc = cell._tc
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = etree.SubElement(tc, qn("w:tcPr"))

    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = etree.SubElement(tcPr, qn("w:tcW"))
    tcW.set(qn("w:w"), str(width))
    tcW.set(qn("w:type"), "dxa")


def _set_vertical_align(cell, align: str = "center"):
    """设置单元格垂直对齐。"""
    tc = cell._tc
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = etree.SubElement(tc, qn("w:tcPr"))

    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = etree.SubElement(tcPr, qn("w:vAlign"))
    vAlign.set(qn("w:val"), align)


def _set_row_height(row, height: int, rule: str = "atLeast"):
    """设置行高。"""
    tr = row._tr
    trPr = tr.find(qn("w:trPr"))
    if trPr is None:
        trPr = etree.SubElement(tr, qn("w:trPr"))

    trHeight = trPr.find(qn("w:trHeight"))
    if trHeight is None:
        trHeight = etree.SubElement(trPr, qn("w:trHeight"))
    trHeight.set(qn("w:val"), str(height))
    trHeight.set(qn("w:hRule"), rule)


def _set_borders(table):
    """设置表格单线边框。"""
    from docx.oxml import parse_xml
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblBorders = tblPr.find(qn("w:tblBorders"))
    if tblBorders is not None:
        tblPr.remove(tblBorders)

    borders_xml = (
        f'<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:top w:val="single" w:sz="{BORDER_SIZE}" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:left w:val="single" w:sz="{BORDER_SIZE}" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:bottom w:val="single" w:sz="{BORDER_SIZE}" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:right w:val="single" w:sz="{BORDER_SIZE}" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:insideH w:val="single" w:sz="{BORDER_SIZE}" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:insideV w:val="single" w:sz="{BORDER_SIZE}" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(parse_xml(borders_xml))


def _create_cell_text(cell, text: str, bold: bool = False, center: bool = False,
                       font_name: str = FONT_NAME, size: str = FONT_SIZE_BODY):
    """在单元格中创建格式化的文本段落。"""
    cell.text = ""
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.clear()
    _clear_para_indent(p)

    # 清除段落默认间距，确保垂直居中真正生效
    pPr = p._p.find(qn("w:pPr"))
    if pPr is None:
        pPr = etree.SubElement(p._p, qn("w:pPr"))
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = etree.SubElement(pPr, qn("w:spacing"))
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")

    if center:
        p.alignment = 1  # CENTER

    run = p.add_run(text)
    _set_font(run, font_name=font_name, size=size, bold=bold)
