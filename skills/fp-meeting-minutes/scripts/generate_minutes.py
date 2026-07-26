#!/usr/bin/env python3
"""
会议纪要 Word 文档生成脚本。

用法:
    python generate_minutes.py --input data.json --output meeting.docx

或作为模块导入:
    from generate_minutes import generate_fp_minutes, generate_solution_minutes
    generate_fp_minutes(data, output_path)

格式参数遵循 references/docx_format_spec.md，底层 XML 操作由 shared_utils.py 封装。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

# 导入公共工具库
sys.path.insert(0, str(Path(__file__).parent))
from shared_utils import (
    FONT_SIZE_BODY,
    FONT_SIZE_H1,
    FONT_SIZE_H2,
    FONT_SIZE_TITLE,
    _clear_para_indent,
    _create_cell_text,
    _set_borders,
    _set_cell_margins,
    _set_cell_width,
    _set_font,
    _set_row_height,
    _set_table_cell_margins,
    _set_vertical_align,
)

# ── generate_minutes 特有常量 ───────────────────────────────────────
# 会议信息表（4列）
MEETING_GRID = [2000, 3000, 2000, 3000]
MEETING_ROW_HEIGHT = 460    # DXA, atLeast

# 遗留事项表（3列）
TODO_GRID = [1600, 6400, 900]
TODO_ROW_HEIGHT = 460       # DXA, atLeast


# ── 段落构建器 ─────────────────────────────────────────────────────

def _add_title(doc, text: str):
    """添加主标题（居中、加粗、2号字）。"""
    p = doc.add_paragraph()
    _clear_para_indent(p)
    p.alignment = 1  # CENTER
    run = p.add_run(text)
    _set_font(run, size=FONT_SIZE_TITLE, bold=True)


def _add_subtitle(doc, text: str):
    """添加副标题（左对齐）。"""
    p = doc.add_paragraph()
    _clear_para_indent(p)
    run = p.add_run(text)
    _set_font(run)


def _add_heading1(doc, text: str):
    """添加一级标题。"""
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    _clear_para_indent(p)
    run = p.add_run(text)
    _set_font(run, size=FONT_SIZE_H1, bold=True)


def _add_heading2(doc, text: str):
    """添加二级标题。"""
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    _clear_para_indent(p)
    run = p.add_run(text)
    _set_font(run, size=FONT_SIZE_H2, bold=True)


def _add_normal_para(doc, text: str):
    """添加正文段落。"""
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    _clear_para_indent(p)
    run = p.add_run(text)
    _set_font(run)


# ── 表格构建器 ─────────────────────────────────────────────────────

def _build_meeting_info_table(doc, data: dict) -> None:
    """构建排产评审会议信息表格（4列）。"""
    table = doc.add_table(rows=0, cols=4)
    _set_borders(table)
    _set_table_cell_margins(table)

    for i, w in enumerate(MEETING_GRID):
        table.columns[i].width = Pt(w / 20)
        gridCol = table._tbl.tblGrid.gridCol_lst[i]
        gridCol.set(qn("w:w"), str(w))

    def _add_row(cols_data: list, merge_spans: list[int] | None = None, bold_labels: bool = True):
        cells = table.add_row().cells
        for ci, cell in enumerate(cells):
            _set_cell_margins(cell)
            _set_vertical_align(cell, "center")
            _set_cell_width(cell, MEETING_GRID[ci])

        for ci, val in enumerate(cols_data):
            if val is None:
                continue
            is_label = False
            if isinstance(val, tuple):
                val, is_label = val

            _create_cell_text(
                cells[ci], val,
                bold=(ci in [0, 2] and bold_labels) or (ci == 0 and merge_spans),
                center=(ci in [0, 2])
            )

        if merge_spans:
            tcPr = cells[1]._tc.find(qn("w:tcPr"))
            if tcPr is None:
                tcPr = etree.SubElement(cells[1]._tc, qn("w:tcPr"))
            gridSpan = tcPr.find(qn("w:gridSpan"))
            if gridSpan is None:
                gridSpan = etree.SubElement(tcPr, qn("w:gridSpan"))
            gridSpan.set(qn("w:val"), "3")
            tr = cells[1]._tc.getparent()
            for ci in range(2, 4):
                tr.remove(cells[ci]._tc)

        _set_row_height(table.rows[-1], MEETING_ROW_HEIGHT, "atLeast")

    mt = data
    _add_row([("会议名称", True), mt.get("meeting_title", "")], merge_spans=[1, 2, 3])
    _add_row([("会议类别", True), mt.get("meeting_type_badge", "☐重大会议 ☑评审会 ☐例会")], merge_spans=[1, 2, 3])
    _add_row([("会议地点", True), mt.get("meeting_location", ""), ("会议时间", True), mt.get("meeting_time", "")])
    _add_row([("参会人员", True), mt.get("attendees", "")], merge_spans=[1, 2, 3])
    _add_row([("纪要整理", True), mt.get("recorder", ""), ("纪要审核", True), mt.get("reviewer", "")])

    doc.add_paragraph()


# ── 解决方案会议信息表（2列）────────────────────────────────────
SOLUTION_MEETING_GRID = [2000, 8000]


def _build_solution_meeting_info_table(doc, data: dict) -> None:
    """构建解决方案业务月度管理会议信息表格（2列）。"""
    table = doc.add_table(rows=0, cols=2)
    _set_borders(table)
    _set_table_cell_margins(table)

    for i, w in enumerate(SOLUTION_MEETING_GRID):
        gridCol = table._tbl.tblGrid.gridCol_lst[i]
        gridCol.set(qn("w:w"), str(w))

    meeting_info = [
        ("主题", data.get("meeting_title", "")),
        ("时间", data.get("meeting_time", "")),
        ("会议方式", data.get("meeting_location", "")),
        ("召集人", data.get("convenor", "沙建银")),
        ("纪要整理", data.get("recorder", "崔哲")),
        ("与会人", data.get("attendees", "")),
        ("议题", data.get("meeting_topics", "工作事项与管理要求说明\n项目进展、风险及问题汇报")),
    ]

    for label, content in meeting_info:
        cells = table.add_row().cells
        for ci, cell in enumerate(cells):
            _set_cell_margins(cell)
            _set_vertical_align(cell, "center")
            _set_cell_width(cell, SOLUTION_MEETING_GRID[ci])
        _create_cell_text(cells[0], label, bold=True, center=True)
        _create_cell_text(cells[1], content, bold=False, center=False)
        _set_row_height(table.rows[-1], 460, "atLeast")

    doc.add_paragraph()


# ── 排产评审遗留事项表（3列）──────────────────────────────────────────
def _build_todo_table(doc, todos: list[dict]) -> None:
    """构建排产评审遗留事项表格（3列），同一项目合并到一行。"""
    table = doc.add_table(rows=0, cols=3)
    _set_borders(table)
    _set_table_cell_margins(table)

    for i, w in enumerate(TODO_GRID):
        gridCol = table._tbl.tblGrid.gridCol_lst[i]
        gridCol.set(qn("w:w"), str(w))

    headers = ["项目", "遗留问题/待办事项", "责任人"]
    cells = table.add_row().cells
    for ci, cell in enumerate(cells):
        _set_cell_margins(cell)
        _set_vertical_align(cell, "center")
        _set_cell_width(cell, TODO_GRID[ci])
        _create_cell_text(cell, headers[ci], bold=True, center=True)
    _set_row_height(table.rows[-1], TODO_ROW_HEIGHT, "atLeast")

    grouped = {}
    for item in todos:
        proj = item.get("project", "")
        if proj not in grouped:
            grouped[proj] = {"details": [], "owners": set()}
        grouped[proj]["details"].append(item.get("detail", ""))
        for o in item.get("owner", "").split("/"):
            o = o.strip()
            if o:
                grouped[proj]["owners"].add(o)

    for proj, gdata in grouped.items():
        cells = table.add_row().cells
        for ci, cell in enumerate(cells):
            _set_cell_margins(cell)
            _set_vertical_align(cell, "center")
            _set_cell_width(cell, TODO_GRID[ci])

        _create_cell_text(cells[0], proj, center=True)

        details = gdata["details"]
        if len(details) == 1:
            detail_text = details[0]
        else:
            detail_text = "\n".join(f"{i}、{d}" for i, d in enumerate(details, 1))
        _create_cell_text(cells[1], detail_text, center=False)

        owner_text = "、".join(sorted(gdata["owners"]))
        _create_cell_text(cells[2], owner_text, center=True)
        _set_row_height(table.rows[-1], TODO_ROW_HEIGHT, "atLeast")

    doc.add_paragraph()


# ── 解决方案遗留事项表（7列）────────────────────────────────────
SOLUTION_TODO_GRID = [700, 3300, 1100, 1100, 1100, 900, 1700]
SOLUTION_TODO_ROW_HEIGHT = 460


def _build_solution_todo_table(doc, todos: list[dict]) -> None:
    """构建解决方案业务遗留事项表格（7列）。"""
    table = doc.add_table(rows=0, cols=7)
    _set_borders(table)
    _set_table_cell_margins(table)

    for i, w in enumerate(SOLUTION_TODO_GRID):
        gridCol = table._tbl.tblGrid.gridCol_lst[i]
        gridCol.set(qn("w:w"), str(w))

    headers = ["序号", "问题/事项", "责任人", "预计闭环\n时间", "实际闭环\n时间", "当前\n状态", "当前进展"]
    cells = table.add_row().cells
    for ci, cell in enumerate(cells):
        _set_cell_margins(cell)
        _set_vertical_align(cell, "center")
        _set_cell_width(cell, SOLUTION_TODO_GRID[ci])
        _create_cell_text(cell, headers[ci], bold=True, center=True)
    _set_row_height(table.rows[-1], 600, "atLeast")

    for item in todos:
        cells = table.add_row().cells
        for ci, cell in enumerate(cells):
            _set_cell_margins(cell)
            _set_vertical_align(cell, "center")
            _set_cell_width(cell, SOLUTION_TODO_GRID[ci])

        _create_cell_text(cells[0], item.get("no", ""), center=True)
        _create_cell_text(cells[1], item.get("detail", ""), center=False)
        _create_cell_text(cells[2], item.get("owner", ""), center=True)
        _create_cell_text(cells[3], item.get("eta", ""), center=True)
        _create_cell_text(cells[4], item.get("closed", ""), center=True)
        _create_cell_text(cells[5], item.get("status", ""), center=True)
        _create_cell_text(cells[6], item.get("progress", ""), center=False)
        _set_row_height(table.rows[-1], SOLUTION_TODO_ROW_HEIGHT, "atLeast")

    doc.add_paragraph()


def _add_fp_project_title(doc, text: str):
    """排产评审项目汇报标题（加粗正文，非 Heading）。"""
    p = doc.add_paragraph()
    _clear_para_indent(p)
    run = p.add_run(text)
    _set_font(run, size=FONT_SIZE_H2, bold=True)


def _add_solution_title(doc, text: str, size="32"):
    """解决方案月度管理会议纪要标题（居中、加粗）。size 为 half-points 字符串。"""
    p = doc.add_paragraph()
    _clear_para_indent(p)
    p.alignment = 1  # CENTER
    run = p.add_run(text)
    _set_font(run, size=size, bold=True)


def _add_solution_section_title(doc, text: str):
    """解决方案月度管理章节标题（加粗、较大字号）。"""
    p = doc.add_paragraph()
    _clear_para_indent(p)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    _set_font(run, size=FONT_SIZE_H1, bold=True)


def _add_solution_project_title(doc, text: str):
    """解决方案月度管理项目汇报标题（加粗）。"""
    p = doc.add_paragraph()
    _clear_para_indent(p)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    _set_font(run, size=FONT_SIZE_H2, bold=True)


# ── 主生成函数 ─────────────────────────────────────────────────────

def generate_fp_minutes(data: dict, output_path: str) -> str:
    """
    生成排产评审会议纪要 Word 文档。
    格式匹配模板：标题 + 4列信息表 + 项目汇报（1. 项目名（PM：xxx）+（1）...） + 3列遗留表 + 整体关注点
    """
    doc = Document()

    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    _add_title(doc, "会议纪要")
    _add_subtitle(doc, f"会议纪要    {data.get('meeting_type_badge', '☐外部会议 ☑内部会议')}")
    _build_meeting_info_table(doc, data)

    _add_heading1(doc, "一、各项目排产汇报")
    projects = data.get("projects", [])
    if projects:
        for proj in projects:
            _add_fp_project_title(doc, proj.get("name", ""))
            for item in proj.get("items", []):
                _add_normal_para(doc, item)
    else:
        _add_normal_para(doc, "（本次口水稿中未提取到项目排产汇报内容，请确认是否需要补充）")

    _add_heading1(doc, "二、遗留事项/待办")
    _build_todo_table(doc, data.get("todos", []))

    _add_heading1(doc, "三、整体关注点")
    highlights = data.get("highlights", [])
    if highlights:
        for hl in highlights:
            _add_normal_para(doc, hl)
    else:
        _add_normal_para(doc, "（本次口水稿中未提取到整体关注点，请确认是否需要补充）")

    doc.save(output_path)
    return output_path


def generate_solution_minutes(data: dict, output_path: str) -> str:
    """
    生成解决方案业务月度管理会议纪要 Word 文档。
    格式匹配模板：电信与AIoT业务线标题 + 2列信息表 + 重点事项与管理要求说明 + 项目进展（【PM】项目名 + ▸） + 7列遗留表

    data 结构:
    {
        "meeting_title": "2026年7月解决方案业务月度管理会议",
        "meeting_time": "2026-7-10 14:08~15:32",
        "meeting_location": "Welink语音会",
        "convenor": "沙建银",
        "recorder": "崔哲",
        "attendees": "沙建银、崔哲、...",
        "leadership": ["▸ 7月份工时严重不足...", "▸ 开标回款...", ...],
        "projects": [
            {
                "pm": "丁召海",
                "name": "农业项目群（9个执行项目+5个客户线索）",
                "items": ["目前在执行9个项目...", "牧数查（二期）...", ...]
            },
            ...
        ],
        "todos": [
            {
                "no": "1",
                "detail": "农业项目群：7月务必把农业信息采集2026运维项目（90万）签回...",
                "owner": "丁召海",
                "eta": "2026.07.31",
                "closed": "/",
                "status": "进行中",
                "progress": "上领导班子会流程已走完，等待排进班子会议程..."
            },
            ...
        ]
    }
    """
    doc = Document()

    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    # 标题
    _add_solution_title(doc, "电信与AIoT业务线", size="32")
    _add_solution_title(doc, "解决方案项目月度管理会议纪要", size="32")
    doc.add_paragraph()  # 空行

    # 会议信息表（2列）
    _build_solution_meeting_info_table(doc, data)

    # 重点事项与管理要求说明
    leadership = data.get("leadership", [])
    if leadership:
        _add_solution_section_title(doc, "重点事项与管理要求说明")
        for item in leadership:
            _add_normal_para(doc, item)
        doc.add_paragraph()
    else:
        # 如果口水稿中未提取到管理要求，生成提示占位
        _add_solution_section_title(doc, "重点事项与管理要求说明")
        _add_normal_para(doc, "▸ （本次口水稿中未提取到明确的管理要求，请确认是否需要补充）")
        doc.add_paragraph()

    # 项目进展、风险及问题汇报
    projects = data.get("projects", [])
    _add_solution_section_title(doc, "项目进展、风险及问题汇报")
    if projects:
        for proj in projects:
            pm = proj.get("pm", "")
            name = proj.get("name", "")
            _add_solution_project_title(doc, f"【{pm}】{name}")
            for item in proj.get("items", []):
                _add_normal_para(doc, item)
            doc.add_paragraph()
    else:
        _add_normal_para(doc, "（本次口水稿中未提取到项目汇报内容，请确认是否需要补充）")
        doc.add_paragraph()

    # 遗留问题/待办事项跟踪
    todos = data.get("todos", [])
    _add_solution_section_title(doc, "遗留问题/待办事项跟踪")
    _build_solution_todo_table(doc, todos)

    doc.save(output_path)
    return output_path


# ── CLI 入口 ──────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="生成会议纪要 Word 文档")
    parser.add_argument("--input", "-i", required=True, help="JSON 数据文件路径")
    parser.add_argument("--output", "-o", required=True, help="输出 docx 文件路径")
    parser.add_argument("--type", "-t", default="fp", choices=["fp", "solution"],
                        help="会议类型: fp=排产评审, solution=解决方案业务月度管理")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        data = json.load(f)

    if args.type == "fp":
        path = generate_fp_minutes(data, args.output)
    else:
        path = generate_solution_minutes(data, args.output)

    print(f"Document saved to: {path}")


if __name__ == "__main__":
    main()
