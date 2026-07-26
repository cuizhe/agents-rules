#!/usr/bin/env python3
"""
会议纪要 docx 格式修复工具。

自动修复以下常见问题：
1. 标题格式（居中 / 加粗 / 22pt / 微软雅黑）
2. 表格单元格缩进（清除 Normal 样式继承的 firstLineChars）
3. 表格单元格边距清零（tblCellMar + tcMar）

用法:
    python fix_docx_format.py <input.docx> [output.docx]

若不指定 output，则覆盖原文件。
"""

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

# 导入公共工具库
sys.path.insert(0, str(Path(__file__).parent))
from shared_utils import (
    FONT_SIZE_TITLE,
    _clear_para_indent,
    _set_cell_margins,
    _set_font,
    _set_table_cell_margins,
)


def fix_title_format(doc):
    """
    修复最上方两行标题格式：居中、加粗、22pt、微软雅黑。
    假设标题位于文档开头的前两个非空段落。
    """
    title_paras = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            title_paras.append(p)
        if len(title_paras) >= 2:
            break

    for p in title_paras:
        text = p.text.strip()  # 保存原文
        p.clear()
        _clear_para_indent(p)
        p.alignment = 1  # CENTER
        run = p.add_run(text)
        _set_font(run, size=FONT_SIZE_TITLE, bold=True)


def fix_table_indent_and_margins(doc):
    """
    修复所有表格的单元格缩进和边距。
    关键：Normal 样式可能定义 w:firstLineChars="200"，需要显式覆盖。
    """
    for table in doc.tables:
        # 1. 表格级别边距清零
        _set_table_cell_margins(table)

        # 2. 单元格级别边距清零 + 段落缩进清零
        for row in table.rows:
            for cell in row.cells:
                _set_cell_margins(cell)
                for p in cell.paragraphs:
                    _clear_para_indent(p)
                    # 额外清理 firstLineChars（如果 _clear_para_indent 未覆盖到）
                    pPr = p._p.find(qn("w:pPr"))
                    if pPr is not None:
                        ind = pPr.find(qn("w:ind"))
                        if ind is not None:
                            for attr in [qn("w:firstLineChars"), qn("w:firstLine")]:
                                if attr in ind.attrib:
                                    del ind.attrib[attr]


def main():
    if len(sys.argv) < 2:
        print("用法: python fix_docx_format.py <input.docx> [output.docx]")
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else input_path

    print(f"[1/3] 读取文档: {input_path}")
    doc = Document(input_path)

    print(f"[2/3] 修复标题格式...")
    fix_title_format(doc)

    print(f"[3/3] 修复表格缩进与边距...")
    fix_table_indent_and_margins(doc)

    doc.save(output_path)
    print(f"      完成，已保存: {output_path}")


if __name__ == "__main__":
    main()
