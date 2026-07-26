# -*- coding: utf-8 -*-
from openpyxl import load_workbook

wb = load_workbook('references/test_output_v7.xlsx', data_only=False)

# 检查汇总表累计成本行
for sheet_name in ['当月汇总毛利率', '2026年汇总毛利率']:
    ws = wb[sheet_name]
    print(f'=== {sheet_name} 行9（累计成本）===')
    for col_idx in range(1, 10):
        cell = ws.cell(row=9, column=col_idx)
        val_repr = repr(cell.value)[:35] if cell.value is not None else 'None'
        type_str = f"[{cell.data_type}]" if cell.data_type == 'f' else ""
        print(f'  col{col_idx}: {val_repr} {type_str}')
    print()

# 检查明细表第一个项目
ws = wb['当月项目明细数据']
print('=== 当月项目明细数据 - 第一个项目累计成本行（行7）===')
for col_idx in range(18, 25):
    cell = ws.cell(row=7, column=col_idx)
    val_repr = repr(cell.value)[:35] if cell.value is not None else 'None'
    type_str = f"[{cell.data_type}]" if cell.data_type == 'f' else ""
    print(f'  col{col_idx}: {val_repr} {type_str}')

wb.close()
