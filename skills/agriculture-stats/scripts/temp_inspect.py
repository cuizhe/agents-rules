import openpyxl
from pathlib import Path

input_file = list(Path('.').glob('*毛利管控*'))[0]
wb = openpyxl.load_workbook(input_file, read_only=True, data_only=True)
sheet_name = [n for n in wb.sheetnames if '实际' in n][0]
ws = wb[sheet_name]

# 读取所有项目
projects = []
current = None
for row in range(2, ws.max_row + 1):
    code = ws.cell(row=row, column=1).value
    name = ws.cell(row=row, column=2).value
    if code and str(code).strip() not in ('', 'None'):
        if current:
            projects.append(current)
        current = {'code': code, 'name': name, 'start_row': row}
if current:
    projects.append(current)

print(f'总项目数: {len(projects)}')

# 关键字
agri_kw = ['农业','农村','农田','种植','养殖','畜牧','渔业','农机','种业','农资','农产品','畜禽','林业','农场','审计署','兽医','中监所']
exclude_kw = ['银行','保险','证券','贷款','融资']

agri_projects = []
for p in projects:
    name = str(p['name']) if p['name'] else ''
    if any(ex in name for ex in exclude_kw):
        continue
    if any(kw in name for kw in agri_kw):
        agri_projects.append(p)

print(f'农业项目数: {len(agri_projects)}')
for p in agri_projects:
    print(f'  {p["code"]}: {p["name"][:70]}')

# 统计核算科目
print()
print('=== 核算科目（第18列，汇总）===')
categories = set()
for row in range(2, ws.max_row + 1):
    cat = ws.cell(row=row, column=18).value
    if cat and str(cat).strip() not in ('', 'None'):
        categories.add(cat)
for c in sorted(categories):
    print(f'  {c}')

wb.close()