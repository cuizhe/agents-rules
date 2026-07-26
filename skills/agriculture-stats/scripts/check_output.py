import openpyxl
from pathlib import Path

ref_dir = Path(__file__).parent.parent / 'references'
output_file = list(ref_dir.glob('*汇总*.xlsx'))[0]

wb = openpyxl.load_workbook(output_file)
ws_summary = wb['当月汇总毛利率']

print("=== 当月汇总毛利率 - 完整数据 ===")
print("\n列名（第 3 行）:")
headers = {}
for col in range(1, 20):
    val = ws_summary.cell(row=3, column=col).value
    if val:
        headers[col] = str(val)[:15]
        print(f"  {col}: {headers[col]}")

print("\n\n数据（第 4-7 行）:")
for row in range(4, 8):
    indicator = ws_summary.cell(row=row, column=2).value
    print(f"\n{indicator}:")
    
    for col in range(3, 20):
        val = ws_summary.cell(row=row, column=col).value
        if val is not None and val != '':
            header = headers.get(col, f'列{col}')
            print(f"  {header}: {val}")
