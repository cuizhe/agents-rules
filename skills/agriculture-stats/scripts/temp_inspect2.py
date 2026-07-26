import openpyxl
from pathlib import Path

input_file = list(Path('.').glob('*毛利管控*'))[0]
wb = openpyxl.load_workbook(input_file, read_only=True, data_only=True)

print('所有工作表:', wb.sheetnames)
print()

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f'=== {sname} ===')
    # 读取前2行表头
    for row in range(1, 3):
        row_data = [(col, ws.cell(row=row, column=col).value) for col in range(1, 20) if ws.cell(row=row, column=col).value is not None]
        if row_data:
            print(f'  行{row}: {row_data}')
    print()

wb.close()