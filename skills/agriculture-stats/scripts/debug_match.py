import openpyxl
from pathlib import Path

ref_dir = Path(__file__).parent.parent / 'references'
output_file = list(ref_dir.glob('*汇总*.xlsx'))[0]

wb = openpyxl.load_workbook(output_file)
ws_summary = wb['当月汇总毛利率']

# 要查找的指标
test_indicators = [
    '累计收入（NR）',
    '人工成本（成本 1）',
    '差旅费（成本 3）',
    '其他成本（成本 3）'
]

print("=== 汇总表 B 列实际值 ===")
for row in range(4, 8):
    val = ws_summary.cell(row=row, column=2).value
    print(f"行{row}: '{val}' (len={len(val)})")

print("\n=== 逐个字符比较 ===")
for indicator in test_indicators:
    print(f"\n查找：'{indicator}' (len={len(indicator)})")
    
    for row in range(4, 8):
        cell_val = ws_summary.cell(row=row, column=2).value
        if cell_val:
            print(f"  vs 行{row}: '{cell_val}' (len={len(cell_val)})")
            
            # 逐字符比较
            if len(cell_val) == len(indicator):
                all_match = True
                for i in range(len(indicator)):
                    if indicator[i] != cell_val[i]:
                        print(f"    [{i}] 不同：'{indicator[i]}' (U+{ord(indicator[i]):04X}) vs '{cell_val[i]}' (U+{ord(cell_val[i]):04X})")
                        all_match = False
                if all_match:
                    print(f"    -> 完全匹配!")
            else:
                print(f"    -> 长度不同，跳过")
            
            # 测试包含关系
            if indicator in str(cell_val):
                print(f"    -> 包含匹配成功!")
