# -*- coding: utf-8 -*-
from openpyxl import load_workbook
import os

out_path = 'references/test_output_v6.xlsx'
wb = load_workbook(out_path, data_only=False)

print('=== 工作表列表 ===')
print(wb.sheetnames)

# 检查汇总表
for sheet_name in ['当月汇总毛利率', '2026年汇总毛利率']:
    if sheet_name not in wb.sheetnames:
        continue
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} ===')
    for row_idx in range(3, 15):
        row_data = {}
        for col_idx in range(1, 20):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                val_repr = repr(cell.value)[:40]
                if cell.data_type == 'f':
                    val_repr = f"[F]{val_repr}"
                row_data[col_idx] = val_repr
        if row_data:
            print(f'  行{row_idx}: {row_data}')

    print('\n=== 合并单元格 ===')
    for m in ws.merged_cells.ranges:
        print(f'  {m}')

# 检查明细表第一个项目
ws = wb['当月项目明细数据']
print('\n=== 当月项目明细数据 - 第一个项目（前12行）===')
for row_idx in range(2, 14):
    row_data = {}
    for col_idx in [1, 2, 3, 4, 5, 18, 19, 20]:
        cell = ws.cell(row=row_idx, column=col_idx)
        if cell.value is not None:
            val_repr = repr(cell.value)[:35]
            if cell.data_type == 'f':
                val_repr = f"[F]{val_repr}"
            row_data[col_idx] = val_repr
    if row_data:
        print(f'  行{row_idx}: {row_data}')

wb.close()
