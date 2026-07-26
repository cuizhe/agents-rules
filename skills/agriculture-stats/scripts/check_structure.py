import openpyxl
from pathlib import Path

ref_dir = Path(__file__).parent.parent / 'references'
output_file = list(ref_dir.glob('*汇总*.xlsx'))[0]

wb = openpyxl.load_workbook(output_file)
ws_summary = wb['当月汇总毛利率']

print("=== 当月汇总毛利率 - 第 1-3 行完整结构 ===")
for row in range(1, 4):
    print(f"\n行{row}:")
    for col in range(1, 25):
        val = ws_summary.cell(row=row, column=col).value
        if val:
            print(f"  列{col} ({chr(64+col) if col <= 26 else chr(64+col//26)+chr(64+col%26)}): {val}")

# 查看明细表的列名
ws_detail = wb['当月项目明细数据']
print("\n=== 当月项目明细数据 - 第 1 行列名 ===")
for col in range(1, 60):
    val = ws_detail.cell(row=1, column=col).value
    if val:
        print(f"  列{col}: {val}")
