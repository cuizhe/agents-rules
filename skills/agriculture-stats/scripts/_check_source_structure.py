# -*- coding: utf-8 -*-
from openpyxl import load_workbook

# 用 data_only=True 读取计算值
wb = load_workbook('C:/Users/Administrator/.workbuddy/skills/agriculture-stats/references/2026年电信与AIoT业务线FP项目毛利管控表20260413.xlsx', data_only=True)
ws = wb['实际运营数据-崔哲']

print('=== P202504000428 完整10行数据（data_only=True）===')
print('项目起始行: 362\n')

for offset in range(10):
    row = 362 + offset
    cat = ws.cell(row=row, column=18).value
    print(f'\n行{row} ({cat}):')
    
    # 关键列
    cols_to_check = [
        (19, '预算'),
        (20, '实际累计'),
        (21, '25年实际'),
        (22, 'YTD12'),
        (23, 'YTD11'),
        (24, 'YTD10'),
        (50, '月度03'),
        (51, '月度02'),
        (52, '月度01'),
    ]
    for col, name in cols_to_check:
        v = ws.cell(row=row, column=col).value
        print(f'  col{col}({name}): {v}')

wb.close()
