# -*- coding: utf-8 -*-
from openpyxl import load_workbook

wb = load_workbook('references/test_output_v5.xlsx', data_only=False)

# 检查汇总表公式和合并单元格
for sheet_name in ['当月汇总毛利率', '2026年汇总毛利率']:
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} ===')
    for row_idx in range(3, 15):
        row_data = {}
        for col_idx in range(1, 20):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                row_data[col_idx] = repr(cell.value)[:35]
        if row_data:
            print(f'  行{row_idx}: {row_data}')
    print('  Merged cells:')
    for m in ws.merged_cells.ranges:
        print(f'    {m}')

# 检查明细表是否有闲置成本行
ws = wb['当月项目明细数据']
print('\n=== 当月项目明细数据 第一个项目块 col18(类别) ===')
for row_idx in range(2, 14):
    cat = ws.cell(row=row_idx, column=18).value
    print(f'  行{row_idx}: {cat}')
