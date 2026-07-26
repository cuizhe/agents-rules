# -*- coding: utf-8 -*-
from openpyxl import load_workbook

wb = load_workbook('C:/Users/Administrator/.workbuddy/skills/agriculture-stats/references/2026年电信与AIoT业务线FP项目毛利管控表20260413.xlsx', data_only=True)
ws = wb['实际运营数据-崔哲']

# P202504000428 起始行362，闲置成本=364
print('=== P202504000428 闲置成本行(364) - data_only=True（计算值）===')
for c in range(1, 60):
    v = ws.cell(row=364, column=c).value
    if v is not None and v != '':
        print(f'  col{c}: {v}')

print('\n=== 对比：累计收入行(362) 有数据的列 ===')
for c in range(1, 60):
    v = ws.cell(row=362, column=c).value
    if v is not None and v != '' and not isinstance(v, str):
        print(f'  col{c}: {v}')
