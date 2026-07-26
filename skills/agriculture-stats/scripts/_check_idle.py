# -*- coding: utf-8 -*-
from openpyxl import load_workbook

# 检查源表闲置成本数据
wb = load_workbook('C:/Users/Administrator/.workbuddy/skills/agriculture-stats/references/2026年电信与AIoT业务线FP项目毛利管控表20260413.xlsx', data_only=True)
ws = wb['实际运营数据-崔哲']

print('=== 源表第一个项目（P202212000496）闲置成本行 ===')
# 第一个项目起始行是2，闲置成本在第4行（offset=2）
for row_idx in range(2, 12):
    cat = ws.cell(row=row_idx, column=18).value
    if cat and '闲置' in str(cat):
        print(f'  行{row_idx}: 类别={cat}')
        for col_idx in [19, 20, 21]:
            val = ws.cell(row=row_idx, column=col_idx).value
            print(f'    col{col_idx} = {val}')

print('\n=== 源表所有项目的闲置成本行数据 ===')
row = 2
while row <= ws.max_row:
    code = ws.cell(row=row, column=1).value
    name = ws.cell(row=row, column=2).value
    if code and name and str(code).strip() not in ('', 'None'):
        # 闲置成本行 = 项目起始行 + 2
        idle_row = row + 2
        idle_cat = ws.cell(row=idle_row, column=18).value
        idle_budget = ws.cell(row=idle_row, column=19).value
        idle_actual = ws.cell(row=idle_row, column=20).value
        print(f'  [{code}] {str(name)[:30]}: 预算={idle_budget}, 实际={idle_actual}')
        row += 10
    else:
        row += 1

wb.close()

# 检查汇总表
print('\n=== 汇总表闲置成本行 ===')
wb_out = load_workbook('C:/Users/Administrator/.workbuddy/skills/agriculture-stats/references/农业项目群预算毛利表汇总表-26.4.17.xlsx', data_only=False)
for sheet_name in ['当月汇总毛利率', '2026年汇总毛利率']:
    ws = wb_out[sheet_name]
    print(f'\n{sheet_name}:')
    for row_idx in range(4, 14):
        cat = ws.cell(row=row_idx, column=2).value
        if cat and '闲置' in str(cat):
            budget = ws.cell(row=row_idx, column=3).value
            actual = ws.cell(row=row_idx, column=4).value
            print(f'  行{row_idx}: 类别={cat}, 预算={budget}, 实际={actual}')

wb_out.close()
