import openpyxl
from pathlib import Path

ref_dir = Path(__file__).parent.parent / 'references'
output_file = list(ref_dir.glob('*汇总*.xlsx'))[0]

wb = openpyxl.load_workbook(output_file)
ws_detail = wb['当月项目明细数据']

# 找到 202602 列
col_202602 = 51  # 根据之前的输出

# 检查所有核算方式在 202602 列的数据
print("=== 检查 202602 列（第 51 列）所有非空值 ===")
for row in range(2, ws_detail.max_row + 1):
    method = ws_detail.cell(row=row, column=18).value
    val = ws_detail.cell(row=row, column=col_202602).value
    
    if val is not None and val != '' and method:
        print(f"行{row}: 核算方式='{method}', 值={val} (类型：{type(val).__name__})")
