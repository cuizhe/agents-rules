import openpyxl
from pathlib import Path

ref_dir = Path(__file__).parent.parent / 'references'
output_file = list(ref_dir.glob('*汇总*.xlsx'))[0]

wb = openpyxl.load_workbook(output_file)
ws_summary = wb['当月汇总毛利率']

# 直接从 Excel 读取第 4-7 行 B 列的值
print("=== 直接从 Excel 读取指标名称 ===")
accounting_methods_from_excel = []
for row in range(4, 8):
    val = ws_summary.cell(row=row, column=2).value
    accounting_methods_from_excel.append(val)
    print(f"行{row}: '{val}' (len={len(val)})")

# 测试查找
print("\n=== 测试查找 ===")
for i, method in enumerate(accounting_methods_from_excel, start=4):
    found_row = None
    for row in range(1, 15):
        cell_val = ws_summary.cell(row=row, column=2).value
        if cell_val == method:
            found_row = row
            break
    
    if found_row:
        print(f"[OK] 行{i}的'{method}' -> 找到于 行{found_row}")
    else:
        print(f"[FAIL] 行{i}的'{method}' -> 未找到")
