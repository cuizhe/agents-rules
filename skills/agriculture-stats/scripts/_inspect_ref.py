# -*- coding: utf-8 -*-
from openpyxl import load_workbook
import os

ref_path = os.path.join('references', '农业项目群预算毛利表汇总表-26.4.17.xlsx')
print('Exists:', os.path.exists(ref_path))

wb = load_workbook(ref_path, data_only=False)
print('Sheets:', wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} (rows={ws.max_row}, cols={ws.max_column}) ===')
    # 表头
    if '明细' in sheet_name:
        for row_idx in range(1, 3):
            row_data = {}
            for col_idx in range(1, min(57, ws.max_column+1)):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    row_data[col_idx] = repr(cell.value)[:40]
            if row_data:
                print(f'  行{row_idx}: {row_data}')
        # 前12行
        for row_idx in range(2, 14):
            row_data = {}
            for col_idx in [1, 2, 18, 19, 20]:
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    row_data[col_idx] = repr(cell.value)[:40]
            if row_data:
                print(f'  行{row_idx}: {row_data}')
    elif '汇总' in sheet_name:
        for row_idx in range(1, 15):
            row_data = {}
            for col_idx in range(1, 20):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    row_data[col_idx] = repr(cell.value)[:40]
            if row_data:
                print(f'  行{row_idx}: {row_data}')
        print('  Merged cells:')
        for m in ws.merged_cells.ranges:
            print(f'    {m}')
