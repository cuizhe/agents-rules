# -*- coding: utf-8 -*-
from openpyxl import load_workbook

# 用 data_only=False 检查源表第一个项目完整结构
wb = load_workbook('C:/Users/Administrator/.workbuddy/skills/agriculture-stats/references/2026年电信与AIoT业务线FP项目毛利管控表20260413.xlsx', data_only=False)
ws = wb['实际运营数据-崔哲']

print('=== 第一个项目 P202212000496 完整10行数据 ===')
for row_idx in range(2, 12):
    cat = ws.cell(row=row_idx, column=18).value
    budget = ws.cell(row=row_idx, column=19).value
    actual = ws.cell(row=row_idx, column=20).value
    
    # 检查该行的数据类型
    budget_type = type(budget).__name__
    actual_type = type(actual).__name__
    
    print(f'行{row_idx}: 类别={cat}')
    print(f'  col19(预算)={budget} [{budget_type}], col20(实际)={actual} [{actual_type}]')
    
    # 如果是公式，显示公式内容
    if ws.cell(row=row_idx, column=19).data_type == 'f':
        print(f'  col19公式: {budget}')
    if ws.cell(row=row_idx, column=20).data_type == 'f':
        print(f'  col20公式: {actual}')

wb.close()
