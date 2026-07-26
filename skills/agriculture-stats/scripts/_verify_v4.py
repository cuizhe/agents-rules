# -*- coding: utf-8 -*-
from openpyxl import load_workbook

wb = load_workbook('references/test_output_v4.xlsx', data_only=False)

# 检查明细表第一个项目块
for sheet_name in ['当月项目明细数据', '2026年项目明细数据']:
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} 前12行 col18(类别) ===')
    for row_idx in range(2, 12):
        cat = ws.cell(row=row_idx, column=18).value
        code = ws.cell(row=row_idx, column=1).value
        name = ws.cell(row=row_idx, column=2).value
        budget = ws.cell(row=row_idx, column=19).value
        actual = ws.cell(row=row_idx, column=20).value
        print(f'  行{row_idx}: cat={repr(cat)[:20]}, code={repr(code)[:15] if code else None}, budget={repr(budget)[:15]}, actual={repr(actual)[:15]}')

# 检查汇总表
for sheet_name in ['当月汇总毛利率', '2026年汇总毛利率']:
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} ===')
    for row_idx in range(3, 15):
        row_data = {}
        for col_idx in range(1, 20):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                row_data[col_idx] = repr(cell.value)[:35]
        if row_data:
            print(f'  行{row_idx}: {row_data}')
    print('  Merged cells:')
    for m in ws.merged_cells.ranges:
        print(f'    {m}')
