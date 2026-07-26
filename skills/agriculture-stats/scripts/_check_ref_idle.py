# -*- coding: utf-8 -*-
from openpyxl import load_workbook

wb = load_workbook('C:/Users/Administrator/.workbuddy/skills/agriculture-stats/references/农业项目群预算毛利表汇总表-26.4.17.xlsx', data_only=False)

for sheet_name in ['当月汇总毛利率', '2026年汇总毛利率']:
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} 闲置成本行(行6) ===')
    for col_idx in range(1, 20):
        cell = ws.cell(row=6, column=col_idx)
        val = cell.value
        dt = cell.data_type
        if val is not None:
            print(f'  col{col_idx}: [{dt}] {repr(val)[:30]}')

wb.close()